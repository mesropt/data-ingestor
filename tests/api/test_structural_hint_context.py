"""The structural-hint resolve path keeps its full context (11-06, D-11-22).

`POST /api/structural-hint/resolve` used to call `service.resolve_or_map`
WITHOUT `schema=`, `sheet=` or `strictness=`, and never checked
`result.date_question` -- four confirmed consequences (RESEARCH.md Pitfall
4): the Python-first crosswalk prefill never ran on the hint path, the
escalation line silently disappeared, the vendor pre-fill (which
`/api/confirm` REQUIRES) was never computed, and an ambiguous date column
behind a structural question stranded the field amber forever -- the exact
Confirm dead-end quick task `260712-qgc` fixed on `/api/upload`, re-created
one route over.

Task 1 pins the retention half: the structural-question `UploadEntry` must
keep `schema_name`, `sheet` and `strictness` alive across the question, and
`sheet` must round-trip the `pending_uploads` persistence boundary with the
`.get()` idiom (an old persisted row without the key still rehydrates).

Task 2 pins the resolve half: each of the four consequences gets its own
regression test, plus the Phase 11 load-bearing one -- resolving a hint on a
multi-sheet workbook re-parses the sheet the human chose, never the
workbook's re-ranked winner.

Harness copied from `tests/api/test_date_format_route.py` (the DI-override +
`_mapper` stub + `monkeypatch.setenv` idiom); registry assertions mirror
`tests/api/test_state.py`'s direct-registry idiom. Never a live Claude call.
"""

from __future__ import annotations

import json

from fastapi.testclient import TestClient

from assayingest import service
from assayingest.api.state import UploadEntry, _entry_from_json, _entry_to_json
from assayingest.domain.models import Alias, FieldMapping, MappingProposal
from assayingest.fields.models import Field, FieldSet
from assayingest.parsing.table import RawTable

from .conftest import verified_user

_SCHEMA_NAME = "assay-potency"

#: Every comma is followed by exactly 3 digits -- thousands grouping and a
#: 3-decimal comma display are indistinguishable, so D-14 asks (mirrors
#: `tests/api/test_hint_and_export.py::_write_ambiguous_csv`). Headers spell
#: the canonical field names, so the D-11-17 implicit self-alias covers them.
_AMBIGUOUS_LOCALE_CSV = b"compound_id;value\nA-1;1,234\nA-2;5,678\n"


def _field_set() -> FieldSet:
    return FieldSet(
        name=_SCHEMA_NAME,
        fields=(Field(name="compound_id"), Field(name="value")),
    )


def _client(profile_store, schema_store):
    from assayingest.api.app import app
    from assayingest.api.deps import get_current_user, get_profile_store, get_schema_store

    app.dependency_overrides[get_profile_store] = lambda: profile_store
    app.dependency_overrides[get_schema_store] = lambda: schema_store
    app.dependency_overrides[get_current_user] = lambda: verified_user()
    return TestClient(app)


def _clear():
    from assayingest.api.app import app

    app.dependency_overrides.clear()


def _upload_ambiguous_with_schema_and_sheet(client) -> dict:
    """Upload a file that raises a StructureQuestion, with a `schema_name`
    and an explicit `sheet` supplied -- the exact request shape a sheet-group
    member re-enters this flow with (D-11-22)."""
    return client.post(
        "/api/upload",
        files={"file": ("ambiguous.csv", _AMBIGUOUS_LOCALE_CSV, "text/csv")},
        data={"schema_name": _SCHEMA_NAME, "sheet": "Week 2"},
    ).json()


# --- Task 1: the structural-question entry retains schema/sheet/strictness ---


def test_structural_question_entry_retains_schema_name_sheet_and_strictness(
    profile_store, schema_store
):
    """The resolve route can only pass what the question branch kept: ALL
    THREE of `schema_name`, `sheet`, `strictness` must survive on the
    retained `UploadEntry` (asserted directly against the registry, the
    `tests/api/test_state.py` idiom)."""
    from assayingest.api.state import registry

    service.promote(_field_set(), created_by="seed@example.com", store=schema_store)
    client = _client(profile_store, schema_store)

    body = _upload_ambiguous_with_schema_and_sheet(client)
    entry = registry.get(body["upload_token"])
    _clear()

    assert body["kind"] == "structural_question"
    assert entry is not None
    assert entry.schema_name == _SCHEMA_NAME
    assert entry.sheet == "Week 2"
    assert entry.strictness == "strict"


# --- Task 1: `sheet` round-trips the pending_uploads persistence boundary ----


def _review_ready_entry(sheet: str | None) -> UploadEntry:
    return UploadEntry(
        field_set=FieldSet(fields=(Field(name="compound_id"),)),
        headers_only=False,
        tmp_path=None,
        table=RawTable(headers=["cmpd"], rows=[["NVS-1"]], source_name="tmpxyz.xlsx"),
        provenance="fresh-claude",
        sheet=sheet,
    )


def test_entry_json_round_trip_preserves_the_retained_sheet():
    rehydrated = _entry_from_json(_entry_to_json(_review_ready_entry("Week 2")))

    assert rehydrated.sheet == "Week 2"


def test_an_old_persisted_row_without_the_sheet_key_still_rehydrates():
    """The `.get()` idiom: a row persisted before the `sheet` key existed has
    nothing truthful to offer and must rehydrate rather than destroy a
    curator's mid-review upload on the deploy that ADDED the key."""
    raw = json.loads(_entry_to_json(_review_ready_entry(None)))
    raw.pop("sheet", None)  # simulate a pre-11-06 persisted row

    rehydrated = _entry_from_json(json.dumps(raw))

    assert rehydrated.sheet is None


# --- Task 2: the resolve route runs the SAME resolution the direct path does -


def _counting_mapper(headers_map: dict[str, str]):
    """A propose_mapping_fn mapping each field name straight to a fixed
    source column (never a real Claude call), COUNTING its invocations --
    the prefill test's whole point is that this is never called at all."""
    calls: list[tuple[str, ...]] = []

    def _fn(table, field_set, client=None, *, headers_only=False):
        calls.append(tuple(field_set.field_names))
        return MappingProposal(
            source_columns=list(table.headers),
            field_mappings=[
                FieldMapping(
                    target_field=name, source_column=headers_map.get(name),
                    confidence=1.0, reasoning="fixture", needs_confirmation=False,
                )
                for name in field_set.field_names
            ],
        )

    return _fn, calls


def _resolve(client, token: str, hint: dict):
    return client.post(
        "/api/structural-hint/resolve", json={"upload_token": token, "hint": hint}
    )


def test_resolve_prefills_from_the_crosswalk_and_never_calls_claude(
    monkeypatch, profile_store, schema_store
):
    """Consequence 1 (Pitfall 4): the headers are FULLY covered by the Schema
    crosswalk (D-11-17 implicit self-aliases), so resolving the hint must
    pre-fill every field at confidence 1.0 and pay Claude for NOTHING --
    exactly what the direct `/api/upload` path already does."""
    fn, calls = _counting_mapper({"compound_id": "compound_id", "value": "value"})
    monkeypatch.setattr(service, "propose_mapping", fn)
    monkeypatch.setenv("ANTHROPIC_API_KEY", "test-key")
    service.promote(_field_set(), created_by="seed@example.com", store=schema_store)

    client = _client(profile_store, schema_store)
    body = _upload_ambiguous_with_schema_and_sheet(client)
    assert body["kind"] == "structural_question"

    response = _resolve(client, body["upload_token"], {"decimal_separator": ","})
    _clear()

    assert response.status_code == 200
    resolved = response.json()
    assert resolved["kind"] == "mapping"
    assert calls == []  # zero Claude calls -- the crosswalk covered everything
    for mapping in resolved["field_mappings"]:
        assert mapping["confidence"] == 1.0
        assert mapping["needs_confirmation"] is False


def test_resolve_carries_the_escalation_line(monkeypatch, profile_store, schema_store):
    """Consequence 2: the resolved response names the Python-vs-Claude split,
    exactly as the direct path's Review screen shows it."""
    fn, _calls = _counting_mapper({"compound_id": "compound_id", "value": "value"})
    monkeypatch.setattr(service, "propose_mapping", fn)
    monkeypatch.setenv("ANTHROPIC_API_KEY", "test-key")
    service.promote(_field_set(), created_by="seed@example.com", store=schema_store)

    client = _client(profile_store, schema_store)
    body = _upload_ambiguous_with_schema_and_sheet(client)

    response = _resolve(client, body["upload_token"], {"decimal_separator": ","})
    _clear()

    assert response.status_code == 200
    assert response.json()["escalation"] == {"python": 2, "claude": 0, "total": 2}


def test_resolve_prefills_the_vendor_and_confirm_accepts_it(
    monkeypatch, profile_store, schema_store
):
    """Consequence 3: `/api/confirm` REQUIRES a vendor (422 otherwise), so a
    resolve path that never computes `vendor_memory` leaves the Review screen
    with an empty vendor box after every structural question. The crosswalk
    knows whose format this is -- the response must say so, and a confirm
    with that vendor must land."""
    fn, _calls = _counting_mapper({"compound_id": "Compound Name", "value": "value"})
    monkeypatch.setattr(service, "propose_mapping", fn)
    monkeypatch.setenv("ANTHROPIC_API_KEY", "test-key")
    schema = service.promote(_field_set(), created_by="seed@example.com", store=schema_store)
    schema_store.add_alias(
        schema.id,
        "compound_id",
        Alias(
            vendor="HelixBio", source_column="Compound Name",
            provenance_kind="manual", provenance_actor="seed@example.com",
            created_at="2026-01-01T00:00:00+00:00",
        ),
    )

    client = _client(profile_store, schema_store)
    upload = client.post(
        "/api/upload",
        files={
            "file": ("helix.csv", b"Compound Name;value\nA-1;1,234\nA-2;5,678\n", "text/csv")
        },
        data={"schema_name": _SCHEMA_NAME},
    ).json()
    assert upload["kind"] == "structural_question"

    response = _resolve(client, upload["upload_token"], {"decimal_separator": ","})
    assert response.status_code == 200
    resolved = response.json()
    assert resolved["remembered_vendor"] == "HelixBio"
    assert resolved["remembered_vendor_source"] == "crosswalk"

    confirmed = client.post(
        "/api/confirm",
        json={
            "upload_token": resolved["upload_token"],
            "field_set": _field_set().to_dict(),
            "field_mappings": [
                {
                    "target_field": "compound_id", "source_column": "Compound Name",
                    "confidence": 1.0, "reasoning": "crosswalk",
                    "needs_confirmation": False, "inferred_value": None, "alternatives": [],
                },
                {
                    "target_field": "value", "source_column": "value",
                    "confidence": 1.0, "reasoning": "crosswalk",
                    "needs_confirmation": False, "inferred_value": None, "alternatives": [],
                },
            ],
            "schema_name": _SCHEMA_NAME,
            "vendor": resolved["remembered_vendor"],
        },
    )
    _clear()

    assert confirmed.status_code == 200
    assert confirmed.json()["ready"] is True


#: Mirrors `tests/api/test_date_format_route.py`'s inline ambiguous-date
#: fixture (every date has day<=12 AND month<=12), with a Value column whose
#: every comma is followed by exactly 3 digits -- so a decimal-locale
#: StructureQuestion is raised FIRST, and the date ambiguity sits BEHIND it.
_AMBIGUOUS_DATE_BEHIND_STRUCTURE_CSV = (
    b"Compound Name;Experiment Date;Value\n"
    b"HLX-100;03/11/2025;1,234\n"
    b"HLX-101;03/11/2025;5,678\n"
    b"HLX-102;04/11/2025;9,012\n"
    b"HLX-110;04/11/2025;3,456\n"
    b"HLX-111;05/11/2025;7,890\n"
    b"HLX-112;05/11/2025;2,345\n"
)


def test_ambiguous_date_behind_a_structural_question_raises_the_date_question(
    monkeypatch, profile_store, schema_store
):
    """Consequence 4 -- the Confirm dead-end, pinned closed: an ambiguous
    date column behind a structural question must raise the date question on
    the resolve, NOT come back as a `mapping` whose date field is amber
    forever (Confirm 422s with no way out -- the exact dead-end quick task
    260712-qgc fixed on `/api/upload`). Answering the date question then
    completes normally."""
    field_set = FieldSet(
        fields=(
            Field(name="compound_id"),
            Field(name="assay_date", type="date"),
            Field(name="value"),
        )
    )
    fn, _calls = _counting_mapper(
        {"compound_id": "Compound Name", "assay_date": "Experiment Date", "value": "Value"}
    )
    monkeypatch.setattr(service, "propose_mapping", fn)
    monkeypatch.setenv("ANTHROPIC_API_KEY", "test-key")

    client = _client(profile_store, schema_store)
    upload = client.post(
        "/api/upload",
        files={"file": ("helix.csv", _AMBIGUOUS_DATE_BEHIND_STRUCTURE_CSV, "text/csv")},
        data={"field_set": json.dumps(field_set.to_dict())},
    ).json()
    assert upload["kind"] == "structural_question"

    response = _resolve(client, upload["upload_token"], {"decimal_separator": ","})
    assert response.status_code == 200
    question = response.json()
    assert question["kind"] == "date_question"
    assert question["columns"][0]["target_field"] == "assay_date"

    completed = client.post(
        "/api/date-format/resolve",
        json={
            "upload_token": question["upload_token"],
            "choices": [{"target_field": "assay_date", "order": "day_first"}],
        },
    )
    _clear()

    assert completed.status_code == 200
    body = completed.json()
    assert body["kind"] == "mapping"
    assay_date = next(m for m in body["field_mappings"] if m["target_field"] == "assay_date")
    assert assay_date["needs_confirmation"] is False


def test_resolve_reparses_the_retained_sheet_never_the_reranked_winner(
    monkeypatch, profile_store, schema_store, tmp_path
):
    """T-11-21, the Phase 11 load-bearing case: the human is working on
    'Week 2'. Resolving a structural hint must re-parse THAT sheet --
    without the retained `sheet=`, `parse()` re-ranks the workbook and the
    bigger 'Week 1' wins, silently mapping a different sheet's columns."""
    from openpyxl import Workbook

    from assayingest.api.state import UploadEntry, registry

    workbook = Workbook()
    week_one = workbook.active
    week_one.title = "Week 1"
    week_one.append(["Compound Name", "IC50"])
    for i in range(12):
        week_one.append([f"CPD-{i}", 10.5 + i])
    week_two = workbook.create_sheet("Week 2")
    week_two.append(["Sample", "Result"])
    for i in range(3):
        week_two.append([f"S-{i}", 40 + i])
    xlsx_path = tmp_path / "weeks.xlsx"
    workbook.save(xlsx_path)

    fn, _calls = _counting_mapper({"compound_id": "Sample", "value": "Result"})
    monkeypatch.setattr(service, "propose_mapping", fn)
    monkeypatch.setenv("ANTHROPIC_API_KEY", "test-key")

    token = registry.put(
        UploadEntry(
            field_set=FieldSet(fields=(Field(name="compound_id"), Field(name="value"))),
            headers_only=False, tmp_path=str(xlsx_path),
            source_file_name="weeks.xlsx", sheet="Week 2",
        )
    )
    client = _client(profile_store, schema_store)

    response = _resolve(client, token, {})
    _clear()

    assert response.status_code == 200
    body = response.json()
    assert body["kind"] == "mapping"
    assert body["source_columns"] == ["Sample", "Result"]  # Week 2's headers, not Week 1's


# =============================================================================
# 12-05 Task 2 -- the wire: the human's confirmed layout round-trips.
#
# BOTH answers must END IN A MAPPED DATASET, and neither may re-return the same
# question. An unanswerable loop on the one surface D-12-15's whole principle
# rests on would be the worst possible bug in this phase.
# =============================================================================

import shutil  # noqa: E402
import tempfile  # noqa: E402
from pathlib import Path  # noqa: E402

import pytest  # noqa: E402

from assayingest.api.state import registry  # noqa: E402

_CORPUS = Path(__file__).resolve().parent.parent.parent / "data" / "synthetic" / "lab_corpus"
#: The phase's driving file: `Patient Info` is TWO side-by-side key-value blocks
#: belonging to ONE record (D-12-13) -- labels col 0 / values col 1 AND labels
#: col 3 / values col 4, rows 1..10 -- and un-pivots to 17 columns, 1 row.
CASCADE = _CORPUS / "cascade_allergy_CS-2026-698392.xlsx"

#: The confirmed key-value answer's blocks, as the panel posts them back:
#: CONFIRMATION of Claude's blocks (12-UI-SPEC Discretion 3), never new indices.
_CASCADE_BLOCKS = [
    {"label_column": 0, "value_columns": [1], "first_row": 1, "last_row": 10},
    {"label_column": 3, "value_columns": [4], "first_row": 1, "last_row": 10},
]

_PATIENT_LABELS = [
    "Name", "Medical Record #", "Date of Birth", "Age", "Sex", "Address",
    "City/State/ZIP", "Phone", "Accession #", "Ordering Provider", "Provider NPI",
    "Specimen Type", "Collected", "Received", "Reported", "Priority", "Fasting Status",
]


def _own_copy(source: Path) -> str:
    with tempfile.NamedTemporaryFile(suffix=source.suffix, delete=False) as tmp:
        path = tmp.name
    shutil.copyfile(source, path)
    return path


def _seed_question_entry(source: Path, sheet: str, field_set: FieldSet) -> str:
    """A retained upload sitting on an unanswered structural question -- the
    exact shape `/api/upload` and `/api/sheets/resolve` both leave behind (the
    `test_group_export.py` registry-seeding idiom)."""
    return registry.put(
        UploadEntry(
            field_set=field_set, headers_only=False, tmp_path=_own_copy(source),
            sheet=sheet, source_file_name=source.name,
        )
    )


def _key_value_hint(blocks=None, **overrides) -> dict:
    hint = {
        "layout": {
            "kind": "key_value",
            "confidence": 1.0,
            "reasoning": "confirmed by the curator",
            "key_value_blocks": blocks if blocks is not None else _CASCADE_BLOCKS,
        }
    }
    hint["layout"].update(overrides)
    return hint


# --- ROUND TRIP A: the key-value loop, closed (SHAPE-02) ----------------------


def test_round_trip_a_a_confirmed_key_value_answer_yields_a_mapped_dataset(
    monkeypatch, profile_store, schema_store
):
    """THE SHAPE-02 LOOP, END TO END: the human selects 'Labels down the side'
    on the layout question, the panel posts CONFIRMATION of Claude's blocks,
    and the cascade `Patient Info` sheet -- a sheet the tool REFUSED before
    this phase -- comes back as a real 17-column mapped dataset. Python did the
    un-pivot; Claude never wrote a value (D-12-03/D-12-10)."""
    fn, _calls = _counting_mapper({"compound_id": "Name", "value": "Age"})
    monkeypatch.setattr(service, "propose_mapping", fn)
    monkeypatch.setenv("ANTHROPIC_API_KEY", "test-key")

    token = _seed_question_entry(CASCADE, "Patient Info", _field_set())
    client = _client(profile_store, schema_store)
    response = _resolve(client, token, _key_value_hint())
    _clear()

    assert response.status_code == 200, response.json()
    body = response.json()
    assert body["kind"] == "mapping"  # NOT the same question, returned again
    assert body["source_columns"] == _PATIENT_LABELS  # 17 labels, un-pivoted
    resolved = {m["target_field"]: m["source_column"] for m in body["field_mappings"]}
    assert resolved == {"compound_id": "Name", "value": "Age"}


def test_round_trip_a_the_key_value_answer_does_not_re_return_the_question(
    monkeypatch, profile_store, schema_store
):
    """The loop is CLOSED, asserted as its own fact: no answer leads back to
    the same question. `answerable_by_hint=True` finally means something."""
    monkeypatch.setattr(service, "propose_mapping", _counting_mapper({})[0])
    monkeypatch.setenv("ANTHROPIC_API_KEY", "test-key")

    token = _seed_question_entry(CASCADE, "Patient Info", _field_set())
    client = _client(profile_store, schema_store)
    body = _resolve(client, token, _key_value_hint()).json()
    _clear()

    assert body["kind"] != "structural_question"


# --- ROUND TRIP B: the ordinary-table loop, closed (the REJECT path) ----------


def _ordinary_workbook(tmp_path) -> Path:
    """A sheet the judge could NOT read (its verdict is `unknown`) which is in
    truth a perfectly ordinary table under a banner -- the human rejects the
    tool's non-answer and says so."""
    from openpyxl import Workbook

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Results"
    worksheet.append(["ACME LABS — Batch 7 (synthetic)"])
    worksheet.append(["compound_id", "value"])
    for i in range(1, 6):
        worksheet.append([f"A-{i}", 1.5 + i])
    path = tmp_path / "ordinary.xlsx"
    workbook.save(path)
    return path


@pytest.mark.parametrize(
    "hint, form",
    [
        (
            {"layout": {"kind": "row_per_record", "confidence": 1.0,
                        "reasoning": "confirmed by the curator", "header_row_index": 1}},
            "explicit-layout",
        ),
        ({"header_row_index": 1}, "bare-header_row_index"),
    ],
    ids=["explicit-layout", "bare-header_row_index"],
)
def test_round_trip_b_one_row_per_record_terminates_in_a_mapped_dataset(
    monkeypatch, profile_store, schema_store, tmp_path, hint, form
):
    """THE REJECT PATH, and the one that closes D-12-15's loop: the human
    answers 'One row per record' with the header row, and the sheet maps.

    BOTH answer forms must terminate. The explicit-layout form is what the
    panel posts (12-UI-SPEC Discretion 3) and reads through the VERDICT path.
    The bare-`header_row_index` form is the legacy PARSE-06 hint the number
    input and the clickable grid rows still produce, and it reads through the
    header path. Neither may come back as the SAME question -- that would be an
    unanswerable loop on the exact surface the whole principle rests on.

    The bare form is deliberately pinned HERE, from the plan that depends on
    it: today it terminates via the shape classifier, which Wave C DELETES.
    Plan 12-09's promotion rule (an explicit header row IS a row_per_record
    confirmation) is what must keep it alive -- and if 12-09 forgets, this
    test fails loudly instead of the loop silently reopening."""
    fn, _calls = _counting_mapper({"compound_id": "compound_id", "value": "value"})
    monkeypatch.setattr(service, "propose_mapping", fn)
    monkeypatch.setenv("ANTHROPIC_API_KEY", "test-key")

    token = _seed_question_entry(_ordinary_workbook(tmp_path), "Results", _field_set())
    client = _client(profile_store, schema_store)
    response = _resolve(client, token, hint)
    _clear()

    assert response.status_code == 200, response.json()
    body = response.json()
    assert body["kind"] == "mapping", form  # the loop TERMINATES, both ways
    assert body["source_columns"] == ["compound_id", "value"]  # row 1, not the banner


# --- the wire boundary: a bad index is a 422, never an IndexError -> 500 ------
# (T-12-15, ASVS V5 -- the WR-02 bug class, fixed once, not reintroduced)


@pytest.mark.parametrize(
    "layout, what",
    [
        ({"kind": "sideways", "confidence": 1.0, "reasoning": "x"}, "a kind that does not exist"),
        (
            {"kind": "key_value", "confidence": 1.0, "reasoning": "x",
             "key_value_blocks": [{"label_column": -1, "value_columns": [1],
                                   "first_row": 0, "last_row": 1}]},
            "a negative label column",
        ),
        (
            {"kind": "key_value", "confidence": 1.0, "reasoning": "x",
             "key_value_blocks": [{"label_column": 0, "value_columns": [-2],
                                   "first_row": 0, "last_row": 1}]},
            "a negative value column",
        ),
        (
            {"kind": "key_value", "confidence": 1.0, "reasoning": "x",
             "key_value_blocks": [{"label_column": 0, "value_columns": [1],
                                   "first_row": -5, "last_row": 1}]},
            "a negative first row",
        ),
        (
            {"kind": "row_per_record", "confidence": 1.0, "reasoning": "x",
             "header_row_index": -3},
            "a negative header row",
        ),
        (
            {"kind": "row_per_record", "confidence": 1.0, "reasoning": "x",
             "header_row_index": 0, "first_data_row": 8, "last_data_row": 2},
            "an inverted data-row range",
        ),
        (
            {"kind": "key_value", "confidence": 1.0, "reasoning": "x",
             "key_value_blocks": [{"label_column": 0, "value_columns": [1],
                                   "first_row": 5, "last_row": 1}]},
            "an inverted block row range",
        ),
        ({"kind": "row_per_record", "confidence": 4.2, "reasoning": "x"}, "a confidence above 1"),
    ],
    ids=[
        "bad-kind", "negative-label-column", "negative-value-column", "negative-first-row",
        "negative-header-row", "inverted-data-rows", "inverted-block-rows", "confidence-out-of-range",
    ],
)
def test_a_structurally_impossible_layout_is_422_at_the_wire_boundary(
    profile_store, schema_store, tmp_path, layout, what
):
    """Shape and sign are judged by the wire model itself (the `ReconcileChoiceIn.
    decision` Literal idiom: an invalid value is a 422 AT THE BOUNDARY, never a
    silent mis-apply) -- an index that cannot be TRUE of any grid is refused
    before a file is ever opened."""
    token = _seed_question_entry(_ordinary_workbook(tmp_path), "Results", _field_set())
    client = _client(profile_store, schema_store)
    response = _resolve(client, token, {"layout": layout})
    _clear()

    assert response.status_code == 422, what


def test_an_out_of_grid_index_is_422_naming_the_consequence_never_a_500(
    profile_store, schema_store, tmp_path
):
    """T-12-15: an index that is well-FORMED but lands outside THIS file's grid
    (a tampered or stale client payload) is checked against the re-parsed grid
    and refused with the consequence-first detail -- never an `IndexError`
    dressed as a server error (the WR-02 bug class)."""
    token = _seed_question_entry(_ordinary_workbook(tmp_path), "Results", _field_set())
    client = _client(profile_store, schema_store)
    response = _resolve(
        client, token,
        _key_value_hint(
            blocks=[{"label_column": 999, "value_columns": [1000],
                     "first_row": 0, "last_row": 2}]
        ),
    )
    _clear()

    assert response.status_code == 422
    detail = response.json()["detail"]
    assert "nothing was ingested" in detail.lower()  # the consequence, not the symptom


def test_a_huge_last_row_never_materialises_a_giant_list(
    profile_store, schema_store, tmp_path
):
    """A `last_row` of a billion must be BOUNDED against the real grid before
    any allocation -- a 422, in a moment, never a memory-exhausting slice
    (T-12-15's DoS half)."""
    token = _seed_question_entry(_ordinary_workbook(tmp_path), "Results", _field_set())
    client = _client(profile_store, schema_store)
    response = _resolve(
        client, token,
        _key_value_hint(
            blocks=[{"label_column": 0, "value_columns": [1],
                     "first_row": 0, "last_row": 10**9}]
        ),
    )
    _clear()

    assert response.status_code == 422


def test_a_rejected_layout_leaves_no_temp_file_behind(
    profile_store, schema_store, tmp_path
):
    """The refusal cleans up after itself, exactly as every other error branch
    on this route does (CR-03/P2): the retained temp file's bytes leave disk."""
    token = _seed_question_entry(_ordinary_workbook(tmp_path), "Results", _field_set())
    retained = registry.get(token).tmp_path

    client = _client(profile_store, schema_store)
    response = _resolve(
        client, token,
        _key_value_hint(
            blocks=[{"label_column": 999, "value_columns": [1], "first_row": 0, "last_row": 1}]
        ),
    )
    _clear()

    assert response.status_code == 422
    assert not Path(retained).exists()
