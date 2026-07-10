"""openpyxl NORMAL-mode raw-grid reader — native cell types survive, and
chartsheets are structurally excluded from sheet listing (01-RESEARCH.md
Pattern 6/7, Pitfall 2/3; D-17)."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference

from assayingest.parsing.structure.grid import list_worksheets, read_grid

DATA = Path(__file__).resolve().parent.parent / "data" / "synthetic"


def test_read_grid_zephyr_preserves_native_types():
    rows = read_grid(DATA / "zephyr_bio_ZB-2025.xlsx", "Week 1")
    data_row = rows[5]  # first real record, beneath 3 banner rows + 1 blank
    assert data_row[0] == "ZB-100"
    assert isinstance(data_row[2], float)  # 32.051 -- Pitfall 3 regression guard
    assert isinstance(data_row[3], str)  # "uM"


def test_read_grid_zephyr_banner_rows_are_still_present_in_the_raw_grid():
    # Detection is grid.py's caller's job (header.py) -- read_grid itself must
    # not pre-strip anything, or header scoring would have nothing to score.
    rows = read_grid(DATA / "zephyr_bio_ZB-2025.xlsx", "Week 1")
    assert "ZEPHYR" in str(rows[0][0])
    assert rows[4][0] == "Compound ID"  # the real header, at raw index 4


def test_list_worksheets_excludes_chartsheet(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["label", "val"])
    for i in range(5):
        ws.append([f"r{i}", i * 2])

    chartsheet = wb.create_chartsheet(title="ChartOnly")
    chart = BarChart()
    data = Reference(ws, min_col=2, min_row=1, max_row=6)
    chart.add_data(data, titles_from_data=True)
    chartsheet.add_chart(chart)

    path = tmp_path / "with_chartsheet.xlsx"
    wb.save(path)

    worksheets = list_worksheets(path)
    titles = [ws.title for ws in worksheets]
    assert "ChartOnly" not in titles
    assert titles == ["Data"]
    assert all(hasattr(ws, "iter_rows") for ws in worksheets)


def test_read_grid_missing_path_raises_filenotfound():
    with pytest.raises(FileNotFoundError, match="Cannot ingest"):
        read_grid(DATA / "does_not_exist.xlsx")


def test_read_grid_unsupported_extension_raises_valueerror(tmp_path):
    junk = tmp_path / "notes.txt"
    junk.write_text("hello")
    with pytest.raises(ValueError, match=r"\.csv or \.xlsx"):
        read_grid(junk)


def test_grid_module_never_loads_in_read_only_mode():
    # read_only=True silently strips _charts/_images/_rels (Pitfall 2) --
    # required by plan 03's drawing detection, so it must never appear here.
    source = (
        Path(__file__).resolve().parent.parent
        / "src"
        / "assayingest"
        / "parsing"
        / "structure"
        / "grid.py"
    )
    text = source.read_text()
    assert "read_only=True" not in text


def test_grid_module_imports_no_anthropic():
    source = (
        Path(__file__).resolve().parent.parent
        / "src"
        / "assayingest"
        / "parsing"
        / "structure"
        / "grid.py"
    )
    text = source.read_text()
    import_lines = [
        line.strip()
        for line in text.splitlines()
        if line.strip().startswith("import ") or line.strip().startswith("from ")
    ]
    assert not any("anthropic" in line for line in import_lines)
