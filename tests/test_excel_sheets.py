"""A single Excel workbook can carry many sheets — none may be silently dropped."""

import pytest
from openpyxl import Workbook

from assayingest.cli import resolve_tables
from assayingest.parsing.table import parse_file, sheet_names


@pytest.fixture
def workbook(tmp_path):
    """A two-sheet workbook: 'IC50' and 'Inhibition', plus one empty sheet."""
    wb = Workbook()
    first = wb.active
    first.title = "IC50"
    first.append(["cmpd", "potency", "target_gene"])
    first.append(["NVS-1", "12.5", "EGFR"])
    first.append(["NVS-2", "340", "EGFR"])

    second = wb.create_sheet("Inhibition")
    second.append(["ID", "Inhibition %", "Target"])
    second.append(["CC-1", "82.5", "JAK2"])

    wb.create_sheet("Notes")  # deliberately empty

    path = tmp_path / "multi.xlsx"
    wb.save(path)
    return path


def test_sheet_names_lists_every_sheet(workbook):
    assert sheet_names(workbook) == ["IC50", "Inhibition", "Notes"]


def test_sheet_names_empty_for_csv(tmp_path):
    csv = tmp_path / "x.csv"
    csv.write_text("a,b\n1,2\n")
    assert sheet_names(csv) == []


def test_parse_named_sheet(workbook):
    table = parse_file(workbook, sheet="Inhibition")
    assert table.headers == ["ID", "Inhibition %", "Target"]
    assert table.sheet_name == "Inhibition"
    assert table.rows[0] == ["CC-1", "82.5", "JAK2"]


def test_parse_unknown_sheet_raises_with_available_names(workbook):
    with pytest.raises(ValueError, match="IC50"):
        parse_file(workbook, sheet="Nope")


def test_resolve_tables_returns_every_sheet(workbook):
    tables = resolve_tables(str(workbook))
    assert [t.sheet_name for t in tables] == ["IC50", "Inhibition", "Notes"]
    # The empty 'Notes' sheet is surfaced, not dropped — the CLI reports it.
    assert tables[2].row_count == 0


def test_resolve_tables_can_target_one_sheet(workbook):
    tables = resolve_tables(str(workbook), sheet="IC50")
    assert len(tables) == 1
    assert tables[0].sheet_name == "IC50"


def test_resolve_tables_single_csv(tmp_path):
    csv = tmp_path / "one.csv"
    csv.write_text("a,b\n1,2\n")
    tables = resolve_tables(str(csv))
    assert len(tables) == 1
    assert tables[0].sheet_name is None