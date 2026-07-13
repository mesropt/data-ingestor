"""SHEET-03: every parsed Excel table records the worksheet title it came
from (`RawTable.origin_sheet`), set UNCONDITIONALLY — for a single-sheet
workbook too, not only for a >1-sheet one.

`origin_sheet` is deliberately distinct from `sheet_name`: `sheet_name` is a
DISAMBIGUATION TAG that composes into `RawTable.label` (the text the Claude
prompt and the CLI banner print) and is set only when a workbook has more
than one sheet (`table.py`). Provenance (SHEET-03/D-11-15) must be recorded
on every ingest, so it needs its own field that never touches `label`.

These tests pin, RED-first:
  * a one-sheet workbook: `sheet_name is None` (unchanged) AND
    `origin_sheet == "<the worksheet title>"`;
  * an explicit `sheet="..."` on a multi-sheet workbook: `origin_sheet` is
    that sheet's title;
  * a CSV: `origin_sheet is None`;
  * `RawTable.label` is byte-identical to today's for both a CSV and a
    one-sheet workbook — the value the mapper prompt reads must not shift.
"""

from __future__ import annotations

from openpyxl import Workbook

from assayingest.parsing.table import RawTable, parse, parse_file


def _one_sheet_workbook(tmp_path):
    wb = Workbook()
    sheet = wb.active
    sheet.title = "Potency"
    sheet.append(["cmpd", "value"])
    sheet.append(["NVS-1", "12.5"])
    path = tmp_path / "single.xlsx"
    wb.save(path)
    return path


def _multi_sheet_workbook(tmp_path):
    wb = Workbook()
    first = wb.active
    first.title = "Week 1"
    first.append(["cmpd", "value"])
    first.append(["NVS-1", "12.5"])
    second = wb.create_sheet("Week 2")
    second.append(["cmpd", "value"])
    second.append(["NVS-2", "34.0"])
    path = tmp_path / "weeks.xlsx"
    wb.save(path)
    return path


# --- origin_sheet is set on every Excel parse -------------------------------


def test_one_sheet_workbook_records_origin_sheet_but_leaves_sheet_name_none(tmp_path):
    table = parse_file(_one_sheet_workbook(tmp_path))
    assert table.sheet_name is None  # unchanged: no >1-sheet disambiguation tag
    assert table.origin_sheet == "Potency"


def test_multi_sheet_explicit_sheet_records_that_sheets_title_as_origin(tmp_path):
    table = parse_file(_multi_sheet_workbook(tmp_path), sheet="Week 2")
    assert table.origin_sheet == "Week 2"
    assert table.sheet_name == "Week 2"


def test_csv_leaves_origin_sheet_none(tmp_path):
    csv = tmp_path / "batch.csv"
    csv.write_text("cmpd,value\nNVS-1,12.5\n", encoding="utf-8")
    table = parse_file(csv)
    assert table.origin_sheet is None


# --- label must NOT change (it feeds the Claude prompt and the CLI banner) --


def test_label_unchanged_for_a_one_sheet_workbook(tmp_path):
    table = parse_file(_one_sheet_workbook(tmp_path))
    assert table.label == "single.xlsx"  # no " :: Potency" — sheet_name is None


def test_label_unchanged_for_a_csv(tmp_path):
    csv = tmp_path / "batch.csv"
    csv.write_text("cmpd,value\nNVS-1,12.5\n", encoding="utf-8")
    table = parse_file(csv)
    assert table.label == "batch.csv"


# --- the structural parse() path sets it too --------------------------------


def test_structural_parse_of_one_sheet_workbook_sets_origin_sheet(tmp_path):
    outcome = parse(_one_sheet_workbook(tmp_path))
    assert isinstance(outcome, RawTable)
    assert outcome.sheet_name is None
    assert outcome.origin_sheet == "Potency"
